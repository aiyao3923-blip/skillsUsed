#!/usr/bin/env python
"""Generate Chinese software copyright operation manuals.

Primary backend: python-docx.
Fallback backend: HTML/CSS, with optional WeasyPrint PDF rendering.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import re
from pathlib import Path
from typing import Any, Iterable


DEFAULT_META = {
    "system_name": "软件系统",
    "version": "V1.0",
    "doc_type": "用户操作说明手册",
    "date": "",
    "company": "",
}


def load_spec(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in {".yaml", ".yml"}:
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise SystemExit("YAML input requires PyYAML. Use JSON or install PyYAML in myenv13.") from exc
        data = yaml.safe_load(text)
    else:
        data = json.loads(text)
    if not isinstance(data, dict):
        raise SystemExit("Input spec must be a JSON/YAML object.")
    return data


def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    return [value]


def text_of(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def metadata(spec: dict[str, Any]) -> dict[str, str]:
    meta = dict(DEFAULT_META)
    meta.update({k: text_of(v) for k, v in dict(spec.get("metadata") or {}).items()})
    if not meta["date"]:
        from datetime import date

        today = date.today()
        meta["date"] = f"{today.year}年{today.month:02d}月"
    return meta


def allow_generic_defaults(spec: dict[str, Any]) -> bool:
    return bool(spec.get("_allow_generic_defaults") or spec.get("allow_generic_defaults"))


def missing(label: str) -> str:
    return f"【需补充：{label}】"


def fallback_value(spec: dict[str, Any], value: Any, label: str, generic: str) -> str:
    concrete = text_of(value)
    if concrete:
        return concrete
    return generic if allow_generic_defaults(spec) else missing(label)


def fallback_list(spec: dict[str, Any], value: Any, label: str, generic: list[Any]) -> list[Any]:
    concrete = [item for item in as_list(value) if text_of(item)]
    if concrete:
        return concrete
    return generic if allow_generic_defaults(spec) else [missing(label)]


def document_title(meta: dict[str, str]) -> str:
    return f"{meta['system_name']} {meta['version']} {meta['doc_type']}".strip()


def require_docx():
    try:
        import docx  # noqa: F401
    except ImportError as exc:
        raise SystemExit(
            "python-docx is required for --backend docx. "
            "Install it with: python -m pip install python-docx"
        ) from exc


LOCAL_ARTIFACT_PATTERNS = (
    re.compile(r"(?i)(?:[a-z]:[\\/]|\\\\|file:///)[^\r\n]+"),
    re.compile(
        r"(?i)(?<![\w])[^\s\\/<>:\"|?*]+\."
        r"(?:py|db|sqlite|json|ya?ml|md|txt|log|ini|cfg|toml|xml|png|jpe?g|bmp|gif|tiff?|csv|xlsx?|docx|pdf|exe|bat|cmd|ps1|sh)"
        r"(?![\w])"
    ),
)


def contains_local_artifact(value: str) -> bool:
    return any(pattern.search(value) for pattern in LOCAL_ARTIFACT_PATTERNS)


def public_artifact_literals(spec: dict[str, Any]) -> tuple[str, ...]:
    """Review-scoped public package names; never a bypass for local absolute paths."""
    records = spec.get("public_artifacts", [])
    if not isinstance(records, list):
        raise SystemExit("public_artifacts 必须为逐项核对的对象列表。")
    values: list[str] = []
    for index, record in enumerate(records):
        label = f"public_artifacts[{index}]"
        if not isinstance(record, dict) or set(record) != {"value", "reason", "evidence"}:
            raise SystemExit(label + " 必须且仅包含 value、reason、evidence。")
        if any(not isinstance(record[key], str) or not record[key].strip() for key in record):
            raise SystemExit(label + " 的各字段必须为非空字符串。")
        value = record["value"]
        normalized = value.replace("\\", "/")
        pieces = normalized.split("/")
        if (
            value != value.strip()
            or any(ord(char) < 32 for char in value)
            or any(char in value for char in '<>:"|?*')
            or normalized.startswith("/")
            or any(piece in {"", ".", ".."} for piece in pieces)
            or LOCAL_ARTIFACT_PATTERNS[0].search(value)
            or not LOCAL_ARTIFACT_PATTERNS[1].fullmatch(pieces[-1])
        ):
            raise SystemExit(label + " 只允许明确的交付包相对文件名；禁止绝对路径、通配符和目录穿越。")
        if value in values:
            raise SystemExit(label + " 与已有公开文件名重复。")
        values.append(value)
    return tuple(sorted(values, key=len, reverse=True))


def validate_public_content(
    value: Any, *, context: str = "正文", public_literals: tuple[str, ...] = ()
) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if key in {"sources", "image", "path", "public_artifacts"}:
                continue
            validate_public_content(item, context=context, public_literals=public_literals)
        return
    if isinstance(value, (list, tuple)):
        for item in value:
            validate_public_content(item, context=context, public_literals=public_literals)
        return
    text = text_of(value)
    # Absolute machine paths remain forbidden even if an approved basename appears inside.
    if text and LOCAL_ARTIFACT_PATTERNS[0].search(text):
        raise SystemExit(f"{context}检测到本机文件名或路径，请改写为用户可见的功能描述后重新生成。")
    scrubbed = text
    for literal in public_literals:
        # An allowlisted basename must not approve a longer private path or filename.
        boundary = r"[\w/\\.\-]"
        scrubbed = re.sub(r"(?<!" + boundary + ")" + re.escape(literal) + r"(?!" + boundary + ")", "公开文件", scrubbed)
    if scrubbed and contains_local_artifact(scrubbed):
        raise SystemExit(f"{context}检测到本机文件名或路径，请改写为用户可见的功能描述后重新生成。")


def set_mixed_font(run_or_style: Any, east_asia_font: str, latin_font: str = "Times New Roman") -> None:
    from docx.oxml.ns import qn

    font = run_or_style.font
    font.name = latin_font
    if hasattr(run_or_style, "_element"):
        element = run_or_style._element
    else:
        element = run_or_style._r
    rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia_font)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), latin_font)


def set_style_font(style: Any, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    from docx.shared import Pt, RGBColor

    set_mixed_font(style, font_name)
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    if bold is not None:
        style.font.bold = bold


def get_or_add_style(doc: Any, name: str, base: str | None = None) -> Any:
    from docx.enum.style import WD_STYLE_TYPE

    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base:
            style.base_style = styles[base]
        return style


def configure_paragraph_format(
    style: Any,
    *,
    alignment: Any | None = None,
    first_line_pt: float | None = None,
    left_pt: float | None = None,
    before_pt: float | None = None,
    after_pt: float | None = None,
    line_spacing: float | None = None,
) -> None:
    from docx.shared import Pt

    fmt = style.paragraph_format
    if alignment is not None:
        fmt.alignment = alignment
    if first_line_pt is not None:
        fmt.first_line_indent = Pt(first_line_pt)
    if left_pt is not None:
        fmt.left_indent = Pt(left_pt)
    if before_pt is not None:
        fmt.space_before = Pt(before_pt)
    if after_pt is not None:
        fmt.space_after = Pt(after_pt)
    if line_spacing is not None:
        fmt.line_spacing = line_spacing


def configure_styles(doc: Any) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import RGBColor

    normal = doc.styles["Normal"]
    set_style_font(normal, "宋体", 12, False)
    configure_paragraph_format(
        normal,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_pt=24,
        before_pt=0,
        after_pt=6,
        line_spacing=1.5,
    )

    for style_name, size, before, after in [
        ("Heading 1", 16, 6, 6),
        ("Heading 2", 14, 6, 6),
        ("Heading 3", 12, 3, 3),
    ]:
        style = doc.styles[style_name]
        set_style_font(style, "黑体", size, True)
        configure_paragraph_format(
            style,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_pt=0,
            before_pt=before,
            after_pt=after,
            line_spacing=1.5,
        )

    cover_title = get_or_add_style(doc, "封面标题", "Normal")
    set_style_font(cover_title, "黑体", 26, True)
    configure_paragraph_format(cover_title, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, after_pt=12)

    cover_subtitle = get_or_add_style(doc, "封面副标题", "Normal")
    set_style_font(cover_subtitle, "宋体", 16, False)
    configure_paragraph_format(cover_subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, after_pt=8)

    toc_title = get_or_add_style(doc, "目录标题", "Normal")
    set_style_font(toc_title, "黑体", 16, True)
    configure_paragraph_format(toc_title, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, before_pt=12, after_pt=6)

    for level, left_indent in [(1, 0), (2, 24), (3, 48)]:
        # Word uses built-in lowercase names; custom "TOC 1" is renamed on field update.
        name = f"toc {level}"
        try:
            toc_item = doc.styles[name]
        except KeyError:
            toc_item = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        toc_item.base_style = normal
        set_style_font(toc_item, "宋体", 12, False)
        configure_paragraph_format(
            toc_item,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_pt=0,
            left_pt=left_indent,
            before_pt=0,
            after_pt=0,
            line_spacing=1.0,
        )

    body = get_or_add_style(doc, "正文段落", "Normal")
    set_style_font(body, "宋体", 12, False)
    configure_paragraph_format(
        body,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_pt=24,
        before_pt=0,
        after_pt=6,
        line_spacing=1.5,
    )

    caption = get_or_add_style(doc, "图注", "Normal")
    set_style_font(caption, "宋体", 10.5, False)
    configure_paragraph_format(caption, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, after_pt=6, line_spacing=1.0)

    table_text = get_or_add_style(doc, "表格文字", "Normal")
    set_style_font(table_text, "宋体", 11, False)
    configure_paragraph_format(table_text, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, after_pt=0, line_spacing=1.0)

    header_text = get_or_add_style(doc, "页眉", "Normal")
    set_style_font(header_text, "宋体", 10.5, False)
    header_text.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    configure_paragraph_format(header_text, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_pt=0, after_pt=0, line_spacing=1.0)

    page_number = get_or_add_style(doc, "页码", "Normal")
    set_style_font(page_number, "宋体", 10.5, False)
    configure_paragraph_format(page_number, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_pt=0, after_pt=0, line_spacing=1.0)


def setup_section(section: Any) -> None:
    from docx.shared import Cm

    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def add_field(paragraph: Any, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    instr_run = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    instr_run._r.append(instr_text)

    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    result_run = paragraph.add_run("1")
    set_mixed_font(result_run, "宋体")

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def set_update_fields_on_open(doc: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    else:
        update = existing
    update.set(qn("w:val"), "true")


def set_page_number_start(section: Any, start: int = 1) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def set_paragraph_bottom_border(paragraph: Any, color: str = "A6A6A6", size: int = 4) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def clear_paragraph(paragraph: Any) -> None:
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)


def add_header(section: Any, meta: dict[str, str], *, include_page_numbers: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm

    section.header.is_linked_to_previous = False
    header = section.header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    clear_paragraph(paragraph)
    paragraph.style = "页眉"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_bottom_border(paragraph)
    run = paragraph.add_run(document_title(meta))
    set_mixed_font(run, "宋体")
    run.font.size = None
    if include_page_numbers:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(14.66), alignment=WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("\t")
        add_field(paragraph, "PAGE")
        paragraph.add_run(" / ")
        add_field(paragraph, "SECTIONPAGES")


def add_centered_page_footer(section: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    clear_paragraph(paragraph)
    paragraph.style = "页码"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")


def clear_footer(section: Any) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    clear_paragraph(paragraph)


def add_cover(doc: Any, meta: dict[str, str]) -> None:
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph(style="封面标题")
    title.add_run(meta["system_name"])
    sub = doc.add_paragraph(style="封面副标题")
    sub.add_run(meta["version"])
    doc_type = doc.add_paragraph(style="封面副标题")
    doc_type.add_run(meta["doc_type"])
    for _ in range(8):
        doc.add_paragraph()
    if meta.get("company"):
        company = doc.add_paragraph(style="封面副标题")
        company.add_run(meta["company"])
    date = doc.add_paragraph(style="封面副标题")
    date.add_run(meta["date"])


def add_toc(doc: Any) -> None:
    doc.add_paragraph("目录", style="目录标题")
    paragraph = doc.add_paragraph()
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')


def add_heading(doc: Any, text: str, level: int) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_body_paragraph(doc: Any, text: Any, *, style: str = "正文段落") -> None:
    value = text_of(text)
    if not value:
        return
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(value)
    set_mixed_font(run, "宋体")


def add_items(doc: Any, items: Iterable[Any], prefix: str | None = None) -> None:
    for i, item in enumerate(as_list(list(items)), start=1):
        if isinstance(item, dict):
            text = "；".join(f"{k}：{v}" for k, v in item.items() if text_of(v))
        else:
            text = text_of(item)
        if not text:
            continue
        label = f"{i}. " if prefix == "number" else ("- " if prefix == "bullet" else "")
        add_body_paragraph(doc, label + text)


def cell_set_text(cell: Any, value: Any, *, bold: bool = False, align_left: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "表格文字"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if align_left else WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text_of(value))
    set_mixed_font(run, "宋体")
    run.font.size = Pt(11)
    run.bold = bold


def set_cell_border(cell: Any, **kwargs: tuple[str, int] | None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        value = kwargs.get(edge)
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        if value is None:
            element.set(qn("w:val"), "nil")
            continue
        border_val, size = value
        element.set(qn("w:val"), border_val)
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def apply_three_line_borders(table: Any) -> None:
    if not table.rows:
        return
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=None, left=None, bottom=None, right=None, insideH=None, insideV=None)

    top_15pt = ("single", 12)
    middle_10pt = ("single", 8)
    bottom_15pt = ("single", 12)

    if len(table.rows) == 1:
        for cell in table.rows[0].cells:
            set_cell_border(cell, top=top_15pt, bottom=bottom_15pt, left=None, right=None, insideH=None, insideV=None)
        return

    for cell in table.rows[0].cells:
        set_cell_border(cell, top=top_15pt, bottom=middle_10pt, left=None, right=None, insideH=None, insideV=None)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=bottom_15pt, top=None, left=None, right=None, insideH=None, insideV=None)


def add_three_line_table(doc: Any, headers: list[str], rows: list[list[Any]]) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT

    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, header in enumerate(headers):
        cell_set_text(table.rows[0].cells[idx], header, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, header in enumerate(headers):
            value = row_values[idx] if idx < len(row_values) else ""
            cell_set_text(cells[idx], value, align_left=len(text_of(value)) > 12 or idx == len(headers) - 1)
    apply_three_line_borders(table)
    doc.add_paragraph()


def screenshot_insert_width(path: Path) -> Any:
    from docx.image.image import Image as DocxImage
    from docx.shared import Cm

    try:
        image = DocxImage.from_file(str(path))
        pixel_width = int(image.px_width or 0)
        pixel_height = int(image.px_height or 0)
    except Exception as exc:
        raise SystemExit("截图文件无法读取，请重新截取清晰的真实系统界面后再生成。") from exc
    if pixel_width <= 0 or pixel_height <= 0:
        raise SystemExit("截图尺寸无效，请重新截取清晰的真实系统界面后再生成。")

    conservative_natural_width_cm = pixel_width / 120 * 2.54
    return Cm(min(14.65, conservative_natural_width_cm))


def add_image(doc: Any, image_path: Any, caption: str, base_dir: Path, chapter_no: int, figure_no: int) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    value = text_of(image_path)
    if not value:
        return
    path = Path(value)
    if not path.is_absolute():
        path = base_dir / path
    if not path.exists():
        raise SystemExit("真实系统截图缺失，请补充截图后再生成最终手册。")
    doc.add_picture(str(path), width=screenshot_insert_width(path))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # An inline picture must not inherit the body first-line indent.
    doc.paragraphs[-1].paragraph_format.first_line_indent = Pt(0)
    cap = caption or missing("图片图注")
    doc.add_paragraph(f"图 {chapter_no}-{figure_no} {cap}", style="图注")


def add_named_items_table(doc: Any, items: list[Any], headers: list[str]) -> None:
    rows: list[list[Any]] = []
    for item in items:
        if isinstance(item, dict):
            rows.append([item.get("name", ""), item.get("entry", ""), item.get("format", ""), item.get("description", "")])
        else:
            rows.append([text_of(item), "", "", ""])
    if rows:
        add_three_line_table(doc, headers, rows)


def validate_chapters(chapters: Any) -> None:
    """Validate the opt-in body schema without coercing or dropping content."""
    def fail(location: str, message: str) -> None:
        # Report schema locations, never user-supplied text or internal paths.
        raise SystemExit(f"{location}：{message}")

    def check_fields(value: Any, allowed: set[str], required: set[str], location: str) -> None:
        if not isinstance(value, dict):
            fail(location, "必须为对象。")
        if set(value) - allowed:
            fail(location, "含未知字段；仅允许 " + ", ".join(sorted(allowed)) + "。")
        absent = required - set(value)
        if absent:
            fail(location, "缺少字段 " + ", ".join(sorted(absent)) + "。")

    def check_text(value: Any, location: str) -> None:
        if not isinstance(value, str) or not value.strip():
            fail(location, "必须为非空字符串。")

    block_fields = {
        "paragraph": {"type", "text"},
        "steps": {"type", "items"},
        "bullets": {"type", "items"},
        "table": {"type", "headers", "rows"},
        "image": {"type", "image", "caption"},
    }

    def check_block(block: Any, location: str) -> None:
        if not isinstance(block, dict):
            fail(location, "必须为对象。")
        kind = block.get("type")
        if not isinstance(kind, str) or kind not in block_fields:
            fail(location + ".type", "必须为 paragraph、steps、bullets、table 或 image。")
        check_fields(block, block_fields[kind], block_fields[kind], location)
        if kind == "paragraph":
            check_text(block["text"], location + ".text")
        elif kind in {"steps", "bullets"}:
            items = block["items"]
            if not isinstance(items, list) or not items:
                fail(location + ".items", "必须为非空字符串列表。")
            for index, item in enumerate(items):
                check_text(item, f"{location}.items[{index}]")
        elif kind == "table":
            headers = block["headers"]
            if not isinstance(headers, list) or not headers or any(not isinstance(h, str) for h in headers):
                fail(location + ".headers", "必须为至少含一列的字符串列表。")
            rows = block["rows"]
            if not isinstance(rows, list):
                fail(location + ".rows", "必须为行列表。")
            for row_index, row in enumerate(rows):
                row_location = f"{location}.rows[{row_index}]"
                if not isinstance(row, list):
                    fail(row_location, "必须为标量列表。")
                if len(row) != len(headers):
                    fail(row_location, f"列数必须与 headers 一致（需要 {len(headers)} 列）。")
                for column_index, cell in enumerate(row):
                    if not isinstance(cell, (str, int, float, bool, type(None))):
                        fail(f"{row_location}[{column_index}]", "仅支持字符串、数字、布尔值或 null。")
        else:  # image
            check_text(block["image"], location + ".image")
            check_text(block["caption"], location + ".caption")

    def check_node(node: Any, location: str, level: int) -> None:
        if level > 3:
            fail(location, "标题嵌套最多为 3 级。")
        check_fields(node, {"title", "blocks", "sections"}, {"title"}, location)
        check_text(node["title"], location + ".title")
        blocks = node.get("blocks", [])
        sections = node.get("sections", [])
        if not isinstance(blocks, list):
            fail(location + ".blocks", "必须为列表。")
        if not isinstance(sections, list):
            fail(location + ".sections", "必须为列表。")
        if not blocks and not sections:
            fail(location, "必须包含有效的 blocks 或 sections。")
        for index, block in enumerate(blocks):
            check_block(block, f"{location}.blocks[{index}]")
        for index, section in enumerate(sections):
            check_node(section, f"{location}.sections[{index}]", level + 1)

    if not isinstance(chapters, list) or not chapters:
        fail("chapters", "必须为非空列表。")
    for index, chapter in enumerate(chapters):
        check_node(chapter, f"chapters[{index}]", 1)


def add_chapters(doc: Any, chapters: list[dict[str, Any]], input_base: Path) -> None:
    """Render validated chapters using the existing, unmodified layout helpers."""
    def add_node(node: dict[str, Any], number: tuple[int, ...], figure_no: int) -> int:
        label = ".".join(str(part) for part in number)
        add_heading(doc, f"{label} {node['title'].strip()}", len(number))
        for block in node.get("blocks", []):
            kind = block["type"]
            if kind == "paragraph":
                add_body_paragraph(doc, block["text"])
            elif kind in {"steps", "bullets"}:
                add_items(doc, block["items"], prefix="number" if kind == "steps" else "bullet")
            elif kind == "table":
                add_three_line_table(doc, block["headers"], block["rows"])
            elif kind == "image":
                figure_no += 1
                add_image(doc, block["image"], block["caption"].strip(), input_base, number[0], figure_no)
            else:
                raise SystemExit("chapters 含不支持的块类型，请先校验输入。")
        for index, section in enumerate(node.get("sections", []), start=1):
            figure_no = add_node(section, number + (index,), figure_no)
        return figure_no

    for index, chapter in enumerate(chapters, start=1):
        add_node(chapter, (index,), 0)


def generate_docx(spec: dict[str, Any], output: Path, input_base: Path) -> None:
    require_docx()
    from docx import Document
    from docx.enum.section import WD_SECTION_START

    if "chapters" in spec:
        validate_chapters(spec["chapters"])
    validate_public_content(spec, public_literals=public_artifact_literals(spec))
    meta = metadata(spec)

    doc = Document()
    for section in doc.sections:
        setup_section(section)
    configure_styles(doc)

    add_cover(doc, meta)

    toc_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(toc_section)
    set_page_number_start(toc_section, 1)
    add_header(toc_section, meta, include_page_numbers=False)
    add_centered_page_footer(toc_section)
    add_toc(doc)

    body_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(body_section)
    set_page_number_start(body_section, 1)
    add_header(body_section, meta, include_page_numbers=True)
    clear_footer(body_section)

    if "chapters" in spec:
        add_chapters(doc, spec["chapters"], input_base)
        set_update_fields_on_open(doc)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
        return

    overview = dict(spec.get("overview") or {})
    environment = dict(spec.get("environment") or {})
    interfaces = as_list(spec.get("interfaces"))
    modules = as_list(spec.get("modules") or spec.get("features"))
    data_outputs = as_list(spec.get("data_outputs"))
    faq = as_list(spec.get("faq"))

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 编写目的", 2)
    add_body_paragraph(
        doc,
        fallback_value(
            spec,
            overview.get("purpose"),
            "编写目的",
            f"本文档用于说明{meta['system_name']}的运行环境、功能结构、操作流程和结果输出方式，为软件著作权材料整理、系统演示和日常使用提供统一依据。",
        ),
    )
    add_heading(doc, "1.2 读者对象", 2)
    audience = fallback_list(spec, overview.get("audience"), "读者对象", ["系统管理员", "业务操作人员", "软件测试与维护人员"])
    add_items(doc, audience, prefix="bullet")
    add_heading(doc, "1.3 文档范围", 2)
    add_body_paragraph(doc, fallback_value(spec, overview.get("scope"), "文档范围", "本文档覆盖软件安装启动、界面说明、核心功能操作、输入输出数据、常见问题和注意事项。"))

    add_heading(doc, "2 软件概述", 1)
    add_heading(doc, "2.1 系统定位", 2)
    add_body_paragraph(doc, fallback_value(spec, overview.get("positioning"), "系统定位", f"{meta['system_name']}面向实际业务场景提供数据处理、可视化交互、结果导出和运行管理能力。"))
    add_heading(doc, "2.2 主要功能", 2)
    module_names = [m.get("name", "") if isinstance(m, dict) else m for m in modules]
    add_items(doc, fallback_list(spec, overview.get("features") or module_names, "主要功能", module_names or ["数据导入", "功能处理", "结果导出"]), prefix="number")
    add_heading(doc, "2.3 应用场景", 2)
    add_items(doc, fallback_list(spec, overview.get("scenarios"), "应用场景", ["软件功能演示", "业务流程处理", "结果归档与复核"]), prefix="bullet")

    add_heading(doc, "3 运行环境与安装", 1)
    add_heading(doc, "3.1 硬件环境", 2)
    add_items(doc, fallback_list(spec, environment.get("hardware"), "硬件环境", ["处理器：现代多核 CPU。", "内存：建议 8GB 或以上。", "存储：预留足够空间保存软件、数据和导出结果。"]), prefix="bullet")
    add_heading(doc, "3.2 软件环境", 2)
    add_items(doc, fallback_list(spec, environment.get("software"), "软件环境", ["操作系统：Windows 10/Windows 11。"]), prefix="bullet")
    deps = as_list(environment.get("dependencies"))
    if deps:
        add_heading(doc, "3.3 依赖组件", 2)
        add_items(doc, deps, prefix="bullet")
        install_heading = "3.4 安装与启动"
    else:
        install_heading = "3.3 安装与启动"
    add_heading(doc, install_heading, 2)
    add_items(doc, fallback_list(spec, environment.get("install_steps"), "安装步骤", ["获取软件安装包并解压到指定目录。", "检查配置文件、数据目录和运行权限。"]), prefix="number")
    add_items(doc, fallback_list(spec, environment.get("startup_steps"), "启动步骤", ["双击启动程序或执行启动脚本。", "等待主界面加载完成后开始操作。"]), prefix="number")

    add_heading(doc, "4 系统界面说明", 1)
    if not interfaces:
        add_body_paragraph(doc, fallback_value(spec, None, "系统界面说明", "系统主界面通常由导航区、功能操作区、参数设置区、结果展示区和状态日志区组成。"))
    fig_no = 1
    for idx, item in enumerate(interfaces, start=1):
        if isinstance(item, dict):
            name = text_of(item.get("name"), f"界面 {idx}")
            add_heading(doc, f"4.{idx} {name}", 2)
            add_body_paragraph(doc, fallback_value(spec, item.get("description"), f"{name}界面说明", f"{name}用于承载对应功能的操作入口、状态反馈和结果展示。"))
            add_items(doc, item.get("areas"), prefix="bullet")
            add_image(doc, item.get("image"), item.get("caption") or name, input_base, 4, fig_no)
            fig_no += 1
        else:
            add_heading(doc, f"4.{idx} {text_of(item)}", 2)

    add_heading(doc, "5 功能操作说明", 1)
    if not modules:
        add_body_paragraph(doc, fallback_value(spec, None, "功能模块清单", "本章按照软件核心功能模块说明用户的实际操作流程。"))
    fig_no = 1
    for idx, module in enumerate(modules, start=1):
        if not isinstance(module, dict):
            module = {"name": text_of(module)}
        name = text_of(module.get("name"), f"功能模块 {idx}")
        add_heading(doc, f"5.{idx} {name}", 2)
        add_body_paragraph(doc, "功能说明：" + fallback_value(spec, module.get("description"), f"{name}功能说明", f"{name}用于完成对应业务处理。"))
        if module.get("entry"):
            add_body_paragraph(doc, "入口位置：" + text_of(module.get("entry")))
        steps = as_list(module.get("steps"))
        if steps:
            add_heading(doc, f"5.{idx}.1 操作步骤", 3)
            add_items(doc, steps, prefix="number")
        outputs = as_list(module.get("outputs"))
        if outputs:
            add_heading(doc, f"5.{idx}.2 输出结果", 3)
            add_items(doc, outputs, prefix="bullet")
        notes = as_list(module.get("notes"))
        if notes:
            add_heading(doc, f"5.{idx}.3 注意事项", 3)
            add_items(doc, notes, prefix="bullet")
        add_image(doc, module.get("image"), module.get("caption") or name, input_base, 5, fig_no)
        fig_no += 1

    add_heading(doc, "6 数据管理与结果导出", 1)
    if data_outputs:
        add_named_items_table(doc, data_outputs, ["名称", "操作入口", "格式", "说明"])
    else:
        add_body_paragraph(doc, fallback_value(spec, None, "数据管理与结果导出", "系统运行过程中可根据业务需要生成日志、截图、报表或结构化数据文件。用户应在操作结束后检查输出目录，确认结果文件已正常生成。"))

    add_heading(doc, "7 常见问题与注意事项", 1)
    if faq:
        for idx, item in enumerate(faq, start=1):
            if isinstance(item, dict):
                add_heading(doc, f"7.{idx} {text_of(item.get('question'), f'问题 {idx}')}", 2)
                add_body_paragraph(doc, item.get("answer"))
            else:
                add_body_paragraph(doc, item)
    else:
        if allow_generic_defaults(spec):
            defaults = [
                ("程序无法启动", "检查运行环境、软件目录、启动权限以及依赖文件是否完整。"),
                ("数据无法加载", "检查文件格式、路径权限、文件名字符和数据完整性。"),
                ("导出结果为空", "确认已完成处理流程，并检查输出目录和磁盘空间。"),
            ]
            for idx, (q, a) in enumerate(defaults, start=1):
                add_heading(doc, f"7.{idx} {q}", 2)
                add_body_paragraph(doc, a)
        else:
            add_body_paragraph(doc, missing("常见问题与注意事项"))

    add_heading(doc, "8 总结", 1)
    add_body_paragraph(doc, fallback_value(spec, spec.get("summary"), "总结", f"本文档围绕{meta['system_name']}的运行环境、界面组成、核心功能、数据输出和常见问题进行了说明，可作为软件著作权材料提交、系统演示和用户培训的操作依据。"))

    set_update_fields_on_open(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def esc(value: Any) -> str:
    return html.escape(text_of(value))


def html_list(items: Any, ordered: bool = False) -> str:
    tag = "ol" if ordered else "ul"
    body = "".join(f"<li>{esc(item)}</li>" for item in as_list(items) if text_of(item))
    return f"<{tag}>{body}</{tag}>" if body else ""


def generate_html(spec: dict[str, Any], output: Path, input_base: Path, pdf: Path | None = None) -> None:
    if "chapters" in spec:
        raise SystemExit("HTML 降级草稿不支持 chapters 自定义正文；请使用 --backend docx。")
    validate_public_content(spec, public_literals=public_artifact_literals(spec))
    meta = metadata(spec)
    overview = dict(spec.get("overview") or {})
    environment = dict(spec.get("environment") or {})
    interfaces = as_list(spec.get("interfaces"))
    modules = as_list(spec.get("modules") or spec.get("features"))
    data_outputs = as_list(spec.get("data_outputs"))
    faq = as_list(spec.get("faq"))
    module_names = [m.get("name", "") if isinstance(m, dict) else m for m in modules]
    purpose_text = fallback_value(spec, overview.get("purpose"), "编写目的", "本文档用于说明软件的运行环境、功能结构、操作流程和结果输出方式。")
    audience_items = fallback_list(spec, overview.get("audience"), "读者对象", ["系统管理员", "业务操作人员"])
    scope_text = fallback_value(spec, overview.get("scope"), "文档范围", "本文档覆盖软件安装启动、界面说明、核心功能操作、输入输出数据、常见问题和注意事项。")
    positioning_text = fallback_value(spec, overview.get("positioning"), "系统定位", meta["system_name"] + "面向实际业务场景提供数据处理、可视化交互、结果导出和运行管理能力。")
    feature_items = fallback_list(spec, overview.get("features") or module_names, "主要功能", module_names or ["数据导入", "功能处理", "结果导出"])
    scenario_items = fallback_list(spec, overview.get("scenarios"), "应用场景", ["软件功能演示", "业务流程处理", "结果归档复核"])
    hardware_items = fallback_list(spec, environment.get("hardware"), "硬件环境", ["处理器：现代多核 CPU。", "内存：建议 8GB 或以上。"])
    software_items = fallback_list(spec, environment.get("software"), "软件环境", ["操作系统：Windows 10/Windows 11。"])
    install_items = fallback_list(spec, environment.get("install_steps"), "安装步骤", ["解压软件安装包。", "检查配置文件和运行权限。"])
    startup_items = fallback_list(spec, environment.get("startup_steps"), "启动步骤", ["双击启动程序。", "进入主界面。"])
    summary_text = fallback_value(spec, spec.get("summary"), "总结", "本文档可作为软件著作权材料提交、系统演示和用户培训的操作依据。")

    def img_tag(path_value: Any, caption: str, chapter: int, no: int) -> str:
        value = text_of(path_value)
        if not value:
            return ""
        path = Path(value)
        if not path.is_absolute():
            path = input_base / path
        if not path.exists():
            raise SystemExit("真实系统截图缺失，请补充截图后再生成最终手册。")
        src = path.as_uri()
        return f'<figure><img src="{src}" /><figcaption>图 {chapter}-{no} {esc(caption)}</figcaption></figure>'

    module_html = []
    for idx, module in enumerate(modules, start=1):
        if not isinstance(module, dict):
            module = {"name": text_of(module)}
        name = text_of(module.get("name"), f"功能模块 {idx}")
        module_html.append(
            f"<h2>5.{idx} {esc(name)}</h2>"
            f"<p><strong>功能说明：</strong>{esc(fallback_value(spec, module.get('description'), f'{name}功能说明', name + '用于完成对应业务处理。'))}</p>"
            f"{'<p><strong>入口位置：</strong>' + esc(module.get('entry')) + '</p>' if module.get('entry') else ''}"
            f"<h3>5.{idx}.1 操作步骤</h3>{html_list(fallback_list(spec, module.get('steps'), f'{name}操作步骤', []), True)}"
            f"<h3>5.{idx}.2 输出结果</h3>{html_list(fallback_list(spec, module.get('outputs'), f'{name}输出结果', []))}"
            f"<h3>5.{idx}.3 注意事项</h3>{html_list(fallback_list(spec, module.get('notes'), f'{name}注意事项', []))}"
            f"{img_tag(module.get('image'), module.get('caption') or missing('图片图注'), 5, idx)}"
        )
    if not module_html:
        module_html.append(f"<p>{esc(fallback_value(spec, None, '功能模块清单', '本章按照软件核心功能模块说明用户的实际操作流程。'))}</p>")

    interface_html = []
    for idx, item in enumerate(interfaces, start=1):
        if not isinstance(item, dict):
            item = {"name": text_of(item)}
        name = text_of(item.get("name"), f"界面 {idx}")
        interface_html.append(
            f"<h2>4.{idx} {esc(name)}</h2>"
            f"<p>{esc(fallback_value(spec, item.get('description'), f'{name}界面说明', name + '用于承载对应功能的操作入口、状态反馈和结果展示。'))}</p>"
            f"{html_list(fallback_list(spec, item.get('areas'), f'{name}界面区域', []))}"
            f"{img_tag(item.get('image'), item.get('caption') or missing('图片图注'), 4, idx)}"
        )
    if not interface_html:
        interface_html.append(f"<p>{esc(fallback_value(spec, None, '系统界面说明', '系统主界面通常由导航区、功能操作区、参数设置区、结果展示区和状态日志区组成。'))}</p>")

    output_rows = ""
    for item in data_outputs:
        if isinstance(item, dict):
            output_rows += (
                f"<tr><td>{esc(item.get('name'))}</td><td>{esc(item.get('entry'))}</td>"
                f"<td>{esc(item.get('format'))}</td><td>{esc(item.get('description'))}</td></tr>"
            )
    output_table = (
        '<table class="three-line"><thead><tr><th>名称</th><th>操作入口</th><th>格式</th><th>说明</th></tr></thead>'
        f"<tbody>{output_rows}</tbody></table>"
        if output_rows
        else f"<p>{esc(fallback_value(spec, None, '数据管理与结果导出', '系统运行过程中可根据业务需要生成日志、截图、报表或结构化数据文件。'))}</p>"
    )

    faq_html = ""
    for idx, item in enumerate(faq, start=1):
        if isinstance(item, dict):
            faq_html += f"<h2>7.{idx} {esc(item.get('question') or '问题')}</h2><p>{esc(item.get('answer'))}</p>"
    if not faq_html:
        if allow_generic_defaults(spec):
            faq_html = (
                "<h2>7.1 程序无法启动</h2><p>检查运行环境、软件目录、启动权限以及依赖文件是否完整。</p>"
                "<h2>7.2 数据无法加载</h2><p>检查文件格式、路径权限、文件名字符和数据完整性。</p>"
            )
        else:
            faq_html = f"<p>{esc(missing('常见问题与注意事项'))}</p>"

    html_text = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{esc(document_title(meta))}</title>
<style>
@page {{ size: A4; margin: 2.54cm 3.17cm; @top-left {{ content: "{esc(document_title(meta))}"; font-family: "Times New Roman", SimSun, "宋体", serif; font-size: 10.5pt; color: #7f7f7f; }} @top-right {{ content: counter(page) " / " counter(pages); font-family: "Times New Roman", SimSun, "宋体", serif; font-size: 10.5pt; color: #7f7f7f; }} }}
body {{ font-family: "Times New Roman", SimSun, "宋体", serif; font-size: 12pt; line-height: 1.5; text-align: justify; }}
.cover {{ page-break-after: always; text-align: center; padding-top: 5cm; }}
.cover h1 {{ font-family: "Times New Roman", SimHei, "黑体", sans-serif; font-size: 26pt; }}
.cover p {{ font-size: 16pt; }}
h1 {{ font-family: "Times New Roman", SimHei, "黑体", sans-serif; font-size: 16pt; page-break-before: always; }}
h2 {{ font-family: "Times New Roman", SimHei, "黑体", sans-serif; font-size: 14pt; }}
h3 {{ font-family: "Times New Roman", SimHei, "黑体", sans-serif; font-size: 12pt; }}
p {{ text-indent: 24pt; margin: 0 0 6pt 0; }}
figure {{ text-align: center; margin: 12pt 0; }}
figure img {{ max-width: 14.65cm; height: auto; }}
figcaption {{ font-size: 10.5pt; margin-top: 3pt; }}
table.three-line {{ width: 100%; border-collapse: collapse; border-top: 1.5pt solid #000; border-bottom: 1.5pt solid #000; margin: 10pt 0; }}
table.three-line th {{ border-bottom: 1pt solid #000; font-weight: bold; text-align: center; }}
table.three-line td, table.three-line th {{ padding: 5pt; font-size: 11pt; border-left: 0; border-right: 0; }}
</style>
</head>
<body>
<section class="cover"><h1>{esc(meta["system_name"])}</h1><p>{esc(meta["version"])}</p><p>{esc(meta["doc_type"])}</p><p>{esc(meta["company"])}</p><p>{esc(meta["date"])}</p></section>
<h1>1 引言</h1>
<h2>1.1 编写目的</h2><p>{esc(purpose_text)}</p>
<h2>1.2 读者对象</h2>{html_list(audience_items)}
<h2>1.3 文档范围</h2><p>{esc(scope_text)}</p>
<h1>2 软件概述</h1>
<h2>2.1 系统定位</h2><p>{esc(positioning_text)}</p>
<h2>2.2 主要功能</h2>{html_list(feature_items, True)}
<h2>2.3 应用场景</h2>{html_list(scenario_items)}
<h1>3 运行环境与安装</h1>
<h2>3.1 硬件环境</h2>{html_list(hardware_items)}
<h2>3.2 软件环境</h2>{html_list(software_items)}
<h2>3.3 安装与启动</h2>{html_list(install_items, True)}{html_list(startup_items, True)}
<h1>4 系统界面说明</h1>{''.join(interface_html)}
<h1>5 功能操作说明</h1>{''.join(module_html)}
<h1>6 数据管理与结果导出</h1>{output_table}
<h1>7 常见问题与注意事项</h1>{faq_html}
<h1>8 总结</h1><p>{esc(summary_text)}</p>
</body></html>"""
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html_text, encoding="utf-8")

    if pdf:
        try:
            from weasyprint import HTML  # type: ignore
        except ImportError as exc:
            raise SystemExit("WeasyPrint is required for --pdf. Install weasyprint or omit --pdf.") from exc
        pdf.parent.mkdir(parents=True, exist_ok=True)
        HTML(filename=str(output)).write_pdf(str(pdf))


def sample_spec() -> dict[str, Any]:
    return {
        "metadata": {
            "system_name": "示例智能检测系统",
            "version": "V1.0",
            "doc_type": "用户操作说明手册",
            "date": "2026年07月",
            "company": "示例单位",
        },
        "overview": {
            "purpose": "说明系统的安装部署、功能操作、结果导出和常见问题处理方式。",
            "audience": ["系统管理员", "业务操作人员", "软件测试人员"],
            "positioning": "本系统用于对业务数据进行导入、处理、分析和结果导出。",
            "features": ["数据导入", "智能处理", "结果可视化", "报告导出"],
            "scenarios": ["日常业务处理", "演示汇报", "结果归档复核"],
        },
        "environment": {
            "hardware": ["处理器：现代多核 CPU。", "内存：建议 8GB 或以上。"],
            "software": ["操作系统：Windows 10/Windows 11。"],
            "install_steps": ["解压软件安装包。", "双击启动程序。"],
            "startup_steps": ["进入主界面。", "检查状态栏提示。"],
        },
        "interfaces": [
            {"name": "主界面", "description": "主界面用于展示导航、参数设置、任务执行和结果预览。", "areas": ["导航区", "参数区", "结果区"]}
        ],
        "modules": [
            {
                "name": "数据导入",
                "description": "用于选择并加载待处理数据。",
                "entry": "主界面左侧的数据导入按钮。",
                "steps": ["点击数据导入按钮。", "选择数据文件或文件夹。", "确认加载结果。"],
                "outputs": ["界面显示数据列表。"],
                "notes": ["请确保数据格式符合系统要求。"],
            }
        ],
        "data_outputs": [
            {"name": "处理报告", "entry": "结果导出功能", "format": "PDF/CSV", "description": "记录处理结果和关键指标。"}
        ],
        "faq": [{"question": "无法启动怎么办？", "answer": "检查软件目录、运行权限和依赖文件是否完整。"}],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Chinese soft-copyright operation manuals.")
    parser.add_argument("--input", type=Path, help="JSON/YAML project specification.")
    parser.add_argument("--output", type=Path, help="Output .docx or .html path.")
    parser.add_argument("--backend", choices=["docx", "html"], default="docx")
    parser.add_argument("--pdf", type=Path, help="Optional PDF output when using --backend html and WeasyPrint is installed.")
    parser.add_argument("--write-sample", type=Path, help="Write a sample JSON spec and exit.")
    parser.add_argument("--allow-generic-defaults", action="store_true", help="Use generic filler for missing content. By default missing project facts are marked with 【需补充：...】.")
    args = parser.parse_args()

    if args.write_sample:
        args.write_sample.parent.mkdir(parents=True, exist_ok=True)
        args.write_sample.write_text(json.dumps(sample_spec(), ensure_ascii=False, indent=2), encoding="utf-8")
        return

    if not args.input:
        raise SystemExit("--input is required unless --write-sample is used.")
    spec = load_spec(args.input)
    if args.allow_generic_defaults:
        spec["_allow_generic_defaults"] = True
    output = args.output
    if output is None:
        meta = metadata(spec)
        suffix = ".docx" if args.backend == "docx" else ".html"
        output = Path(f"{meta['system_name']}{meta['version']}{meta['doc_type']}{suffix}")

    input_base = args.input.resolve().parent
    if args.backend == "docx":
        generate_docx(spec, output, input_base)
    else:
        generate_html(spec, output, input_base, args.pdf)


if __name__ == "__main__":
    main()
