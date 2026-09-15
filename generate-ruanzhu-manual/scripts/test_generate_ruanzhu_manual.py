"""Generator regression tests; fixtures are NOT real software manuals/screenshots.

Run with the document environment's Python and stdlib unittest. Temporary files
stay under RUANZHU_MANUAL_TEST_ROOT (default: cwd/skill改进工作/生成器代理验证).
Set RUANZHU_MANUAL_BASELINE to an original generator for strict compatibility tests.
Only the reviewed native-TOC identity and zero-indent picture fixes are normalized;
all other XML nodes and package parts remain covered by the original regression.
"""
from __future__ import annotations

import ast
import copy
import importlib.util
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import xml.etree.ElementTree as ET
import zipfile
import zlib

# Importing either generator must not modify an installed skill's bytecode cache.
sys.dont_write_bytecode = True


def load_generator(filename: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


SCRIPT = Path(__file__).with_name("generate_ruanzhu_manual.py")
generator = load_generator(SCRIPT, "manual_generator_under_test")
TEST_ROOT = Path(os.environ.get(
    "RUANZHU_MANUAL_TEST_ROOT",
    str(Path.cwd() / "skill改进工作" / "生成器代理验证"),
)).resolve()
BASELINE = os.environ.get("RUANZHU_MANUAL_BASELINE")
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
META = {
    "system_name": "生成器内部结构验证",
    "version": "V0.0",
    "doc_type": "内部测试稿（非真实系统手册）",
    "date": "2026年09月",
    "company": "内部 fixture",
}
LEGACY_HEADINGS = [
    "1 引言", "2 软件概述", "3 运行环境与安装", "4 系统界面说明",
    "5 功能操作说明", "6 数据管理与结果导出", "7 常见问题与注意事项", "8 总结",
]


def write_fixture_png(filename: Path, width: int, height: int) -> None:
    """A solid-color internal sizing fixture, not a simulated application UI."""
    def chunk(kind: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload))

    filename.parent.mkdir(parents=True, exist_ok=True)
    pixels = (b"\0" + b"\x51\x9a\xb0" * width) * height
    data = b"\x89PNG\r\n\x1a\n"
    data += chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    data += chunk(b"IDAT", zlib.compress(pixels))
    data += chunk(b"IEND", b"")
    filename.write_bytes(data)


def make_images(base: Path) -> None:
    write_fixture_png(base / "images" / "wide.png", 1200, 600)
    write_fixture_png(base / "images" / "small.png", 240, 480)


def paragraph(text: str = "仅用于验证结构和排版，不代表真实软件功能。") -> dict:
    return {"type": "paragraph", "text": text}


def minimal_spec() -> dict:
    return {"metadata": dict(META), "chapters": [{"title": "业务办理", "blocks": [paragraph()]}]}


def library_spec() -> dict:
    return {
        "metadata": {**META, "system_name": "馆藏流转内部验证"},
        "sources": [r"C:\private\fixtures\README.md", "sources/internal.json"],
        "overview": {"purpose": "旧目的不得追加"},
        "environment": {"install_steps": ["旧安装步骤不得追加"]},
        "interfaces": [{"name": "旧界面不得追加"}],
        "modules": [{"name": "旧模块不得追加"}],
        "features": ["旧功能别名不得追加"],
        "data_outputs": [{"name": "旧导出不得追加"}],
        "faq": [{"question": "旧问题不得追加", "answer": "旧回答不得追加"}],
        "summary": "旧总结不得追加",
        "chapters": [
            {
                "title": "借还办理",
                # Deliberately put sections before blocks: render order is contractual.
                "sections": [
                    {
                        "title": "借出登记",
                        "blocks": [
                            {"type": "steps", "items": ["选择内部测试读者。", "确认内部测试记录。"]},
                            {"type": "bullets", "items": ["保留当前选择。", "核对状态提示。"]},
                        ],
                        "sections": [{
                            "title": "核对状态",
                            "blocks": [
                                paragraph("三级标题下的内部测试说明。"),
                                {"type": "table", "headers": ["项目", "数量", "确认", "比率", "备注"],
                                 "rows": [["待借", 2, True, 1.5, None], ["已还", 0, False, 0.0, ""]]},
                                {"type": "image", "image": "images/small.png", "caption": "状态局部（内部测试图）"},
                            ],
                        }],
                    },
                    {"title": "归还登记", "blocks": [
                        {"type": "image", "image": "images/wide.png", "caption": "归还示意（内部测试图）"},
                    ]},
                ],
                "blocks": [
                    paragraph("借还办理的内部测试说明。"),
                    {"type": "image", "image": "images/wide.png", "caption": "办理总览（内部测试图）"},
                ],
            },
            {"title": "馆藏查询", "sections": [
                {"title": "筛选记录", "blocks": [
                    paragraph("查询记录的内部测试说明。"),
                    {"type": "image", "image": "images/wide.png", "caption": "查询示意（内部测试图）"},
                ]},
            ]},
        ],
    }


def inspection_spec() -> dict:
    return {
        "metadata": {**META, "system_name": "设备巡检内部验证"},
        "sources": [r"C:\private\fixtures\inspection.py"],
        "chapters": [
            {"title": "异常闭环", "sections": [
                {"title": "指派复核", "sections": [
                    {"title": "复核结论", "blocks": [
                        paragraph("异常闭环的内部测试说明。"),
                        {"type": "steps", "items": ["选择内部测试条目。", "记录内部测试结论。"]},
                        {"type": "table", "headers": ["任务", "状态"], "rows": [["复核甲", "待确认"]]},
                        {"type": "image", "image": "images/wide.png", "caption": "闭环示意（内部测试图）"},
                    ]},
                ]},
            ]},
            {"title": "巡检登记", "blocks": [
                paragraph("巡检登记的内部测试说明。"),
                {"type": "bullets", "items": ["检查测试状态。", "保留测试说明。"]},
                {"type": "image", "image": "images/small.png", "caption": "登记局部（内部测试图）"},
            ]},
            {"title": "班次交接", "blocks": [paragraph("班次交接的内部测试说明。")], "sections": []},
        ],
    }


def public_artifact_records() -> list[dict[str, str]]:
    """Literal-matching fixtures only, not evidence that real package files exist."""
    return [
        {"value": value, "reason": f"内部核对说明：{label}，不供读者显示。",
         "evidence": f"C:\\internal-review\\{label}-review.md"}
        for value, label in (("applet.py", "启动入口"), ("config/runtime.yaml", "相对配置"), ("exports/report.csv", "导出文件"))
    ]


def legacy_spec() -> dict:
    spec = generator.sample_spec()
    spec["metadata"] = dict(META)
    spec["overview"]["scope"] = "内部结构验证。"
    spec["environment"]["dependencies"] = ["内部测试组件"]
    spec["interfaces"][0].update(image="images/wide.png", caption="界面内部测试图")
    spec["modules"][0].update(image="images/small.png", caption="功能内部测试图")
    spec["summary"] = "仅为原八章兼容性验证。"
    spec["sources"] = [r"C:\private\original\README.md"]
    return spec


def package_parts(filename: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(filename) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def visible_text(parts: dict[str, bytes]) -> str:
    return "\n".join(
        element.text or ""
        for name, data in parts.items() if name.startswith("word/") and name.endswith(".xml")
        for element in ET.fromstring(data).findall(".//w:t", NS)
    )


def headings(document) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading ")]


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        TEST_ROOT.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="unittest-", dir=TEST_ROOT)
        self.work = Path(self.temporary.name).resolve()
        self.addCleanup(self.cleanup_work)
        make_images(self.work)

    def cleanup_work(self):
        # Verify the resolved target before TemporaryDirectory's recursive removal.
        if self.work.resolve().parent != TEST_ROOT:
            raise AssertionError("Refusing cleanup outside the test workspace")
        self.temporary.cleanup()

    def generate(self, spec: dict, name: str = "manual.docx"):
        from docx import Document
        output = self.work / name
        generator.generate_docx(spec, output, self.work)
        self.assertTrue(output.is_file())
        return output, Document(output)

    def reject(self, spec: dict, message: str = "chapters", *, preflight: bool = True):
        from docx import Document
        output = self.work / "must-not-be-created" / "manual.docx"
        with mock.patch("docx.Document", wraps=Document) as document_factory:
            with self.assertRaises(SystemExit) as caught:
                generator.generate_docx(spec, output, self.work)
            if preflight:
                document_factory.assert_not_called()
        self.assertIn(message, str(caught.exception))
        self.assertFalse(output.exists())
        self.assertFalse(output.parent.exists())
        return str(caught.exception)

    def block_spec(self, block):
        spec = minimal_spec()
        spec["chapters"][0]["blocks"] = [block]
        return spec

    def cli(self, *arguments: str, cwd: Path | None = None):
        return subprocess.run(
            [sys.executable, "-B", "-X", "utf8", str(SCRIPT), *map(str, arguments)],
            cwd=cwd or self.work, capture_output=True, text=True, encoding="utf-8", timeout=30,
        )

    def expected_reviewed_layout_xml(self, name: str, original: bytes) -> str:
        """Apply only the two reviewed changes to ORIGINAL XML, never strip nodes."""
        root = ET.fromstring(original)
        w = "{" + NS["w"] + "}"
        if name == "word/styles.xml":
            for level in (1, 2, 3):
                matches = [style for style in root.findall("w:style", NS)
                           if style.find("w:name", NS).get(w + "val") == f"TOC {level}"]
                self.assertEqual(len(matches), 1, f"Original TOC {level} must be unique")
                style = matches[0]
                self.assertEqual(style.get(w + "type"), "paragraph")
                self.assertEqual(style.get(w + "styleId"), f"TOC{level}")
                self.assertEqual(style.get(w + "customStyle"), "1")
                # Preserve every formatting child/attribute, including TOC fonts/indents.
                style.set(w + "styleId", f"toc{level}")
                style.find("w:name", NS).set(w + "val", f"toc {level}")
                del style.attrib[w + "customStyle"]
        elif name == "word/document.xml":
            for paragraph_node in root.findall(".//w:p", NS):
                if paragraph_node.find(".//wp:inline", NS) is None:
                    continue
                properties = paragraph_node.find("w:pPr", NS)
                self.assertIsNotNone(properties)
                self.assertIsNone(properties.find("w:ind", NS), "Original image indent must be inherited")
                alignment = properties.find("w:jc", NS)
                self.assertIsNotNone(alignment)
                self.assertEqual(alignment.get(w + "val"), "center")
                # Only add explicit firstLine=0; do not alter alignment, drawing or runs.
                indent = ET.Element(w + "ind", {w + "firstLine": "0"})
                properties.insert(list(properties).index(alignment), indent)
        else:
            self.fail("No normalization is permitted for " + name)
        return ET.canonicalize(ET.tostring(root, encoding="unicode"))

    def assert_legacy_parts_match(self, original: dict[str, bytes], actual: dict[str, bytes]):
        self.assertEqual(set(original), set(actual))
        for name, data in original.items():
            if name in {"word/document.xml", "word/styles.xml"}:
                expected = self.expected_reviewed_layout_xml(name, data)
                # Compare the complete XML trees, not selected properties or omitted parts.
                observed = ET.canonicalize(ET.tostring(ET.fromstring(actual[name]), encoding="unicode"))
                self.assertEqual(expected, observed, name)
            else:
                self.assertEqual(data, actual[name], name)

    def test_legacy_eight_chapters_and_missing_facts(self):
        filename, document = self.generate({"metadata": dict(META)})
        self.assertEqual([text for style, text in headings(document) if style == "Heading 1"], LEGACY_HEADINGS)
        text = visible_text(package_parts(filename))
        for label in ("编写目的", "读者对象", "文档范围", "系统定位", "主要功能", "应用场景",
                      "硬件环境", "软件环境", "安装步骤", "启动步骤", "系统界面说明",
                      "功能模块清单", "数据管理与结果导出", "常见问题与注意事项", "总结"):
            self.assertIn(generator.missing(label), text)

    def test_legacy_generic_defaults_stay_opt_in(self):
        for key in ("allow_generic_defaults", "_allow_generic_defaults"):
            with self.subTest(key=key):
                filename, document = self.generate({"metadata": dict(META), key: True})
                self.assertNotIn("【需补充：", visible_text(package_parts(filename)))
                self.assertEqual([text for style, text in headings(document) if style == "Heading 1"], LEGACY_HEADINGS)

    def test_legacy_features_alias(self):
        spec = legacy_spec()
        spec["features"] = spec.pop("modules")
        _, document = self.generate(spec)
        self.assertIn(("Heading 2", "5.1 数据导入"), headings(document))
        self.assertEqual(len(document.inline_shapes), 2)

    @unittest.skipUnless(BASELINE and Path(BASELINE).is_file(), "Set RUANZHU_MANUAL_BASELINE for original-source comparison")
    def test_existing_functions_and_cli_source_unchanged(self):
        old_source = Path(BASELINE).read_text(encoding="utf-8")
        new_source = SCRIPT.read_text(encoding="utf-8")
        old_functions = {n.name: ast.get_source_segment(old_source, n) for n in ast.parse(old_source).body if isinstance(n, ast.FunctionDef)}
        new_functions = {n.name: ast.get_source_segment(new_source, n) for n in ast.parse(new_source).body if isinstance(n, ast.FunctionDef)}
        self.assertEqual(set(new_functions) - set(old_functions), {"validate_chapters", "add_chapters", "public_artifact_literals"})
        body_entry_points = {"generate_docx", "generate_html"}  # Existing chapters-mode changes.
        reviewed_helper_changes = {"configure_styles", "add_image", "validate_public_content"}
        for name, source in old_functions.items():
            if name not in body_entry_points | reviewed_helper_changes:
                with self.subTest(function=name):
                    self.assertEqual(new_functions[name], source)
        changed_helpers = {name for name, source in old_functions.items()
                           if name not in body_entry_points and new_functions.get(name) != source}
        self.assertEqual(changed_helpers, reviewed_helper_changes)

    @unittest.skipUnless(BASELINE and Path(BASELINE).is_file(), "Set RUANZHU_MANUAL_BASELINE for DOCX package comparison")
    def test_legacy_docx_all_zip_parts_equal_original(self):
        original = load_generator(Path(BASELINE), "manual_generator_baseline_docx")
        alias = legacy_spec()
        alias["features"] = alias.pop("modules")
        specs = [
            {"metadata": dict(META)}, generator.sample_spec(), legacy_spec(), alias,
            {"metadata": dict(META), "allow_generic_defaults": True},
            {"metadata": dict(META), "_allow_generic_defaults": True},
        ]
        for index, spec in enumerate(specs):
            with self.subTest(fixture=index):
                old_file = self.work / "original.docx"
                new_file, _ = self.generate(copy.deepcopy(spec))
                original.generate_docx(copy.deepcopy(spec), old_file, self.work)
                old_parts, new_parts = package_parts(old_file), package_parts(new_file)
                self.assert_legacy_parts_match(old_parts, new_parts)

    @unittest.skipUnless(BASELINE and Path(BASELINE).is_file(), "Set RUANZHU_MANUAL_BASELINE for normalization boundary checks")
    def test_legacy_comparison_rejects_unrelated_document_changes(self):
        original = load_generator(Path(BASELINE), "manual_generator_baseline_document_guard")
        old_file = self.work / "original.docx"
        original.generate_docx(legacy_spec(), old_file, self.work)
        new_file, _ = self.generate(legacy_spec())
        before, after = package_parts(old_file), package_parts(new_file)
        self.assert_legacy_parts_match(before, after)
        w = "{" + NS["w"] + "}"
        cases = (
            ("body text", ".//w:t", None, "不允许的正文变化"),
            ("page margin", ".//w:pgMar", w + "left", "0"),
            ("image width", ".//wp:inline/wp:extent", "cx", "1"),
            ("TOC field", ".//w:instrText", None, "NUMPAGES"),
            ("image alignment", "image-alignment", w + "val", "left"),
            ("extra image indent", "image-indent", w + "left", "480"),
        )
        for label, selector, attribute, value in cases:
            with self.subTest(mutation=label):
                root = ET.fromstring(after["word/document.xml"])
                if selector.startswith("image-"):
                    image = next(p for p in root.findall(".//w:p", NS) if p.find(".//wp:inline", NS) is not None)
                    target = image.find("w:pPr/w:jc" if selector == "image-alignment" else "w:pPr/w:ind", NS)
                else:
                    target = root.find(selector, NS)
                self.assertIsNotNone(target)
                if attribute is None:
                    target.text = value
                else:
                    target.set(attribute, value)
                tampered = {**after, "word/document.xml": ET.tostring(root, encoding="utf-8")}
                with self.assertRaises(AssertionError):
                    self.assert_legacy_parts_match(before, tampered)
        with self.subTest(mutation="header field, a non-normalized part"):
            self.assertIn(b"SECTIONPAGES", after["word/header2.xml"])
            tampered = {**after, "word/header2.xml": after["word/header2.xml"].replace(b"SECTIONPAGES", b"NUMPAGES", 1)}
            with self.assertRaises(AssertionError):
                self.assert_legacy_parts_match(before, tampered)

    @unittest.skipUnless(BASELINE and Path(BASELINE).is_file(), "Set RUANZHU_MANUAL_BASELINE for normalization boundary checks")
    def test_legacy_comparison_rejects_unrelated_style_changes(self):
        original = load_generator(Path(BASELINE), "manual_generator_baseline_styles_guard")
        old_file = self.work / "original.docx"
        original.generate_docx(legacy_spec(), old_file, self.work)
        new_file, _ = self.generate(legacy_spec())
        before, after = package_parts(old_file), package_parts(new_file)
        self.assert_legacy_parts_match(before, after)
        w = "{" + NS["w"] + "}"
        toc = ".//w:style[@w:styleId='toc1']"
        cases = (
            ("Normal indentation", ".//w:style[@w:styleId='Normal']/w:pPr/w:ind", w + "firstLine", "0"),
            ("heading font", ".//w:style[@w:styleId='Heading1']/w:rPr/w:rFonts", w + "ascii", "Arial"),
            ("TOC font", toc + "/w:rPr/w:rFonts", w + "ascii", "Arial"),
            ("TOC left indent", toc + "/w:pPr/w:ind", w + "left", "480"),
            ("TOC first-line indent", toc + "/w:pPr/w:ind", w + "firstLine", "480"),
            ("TOC custom flag", toc, w + "customStyle", "1"),
            ("TOC uppercase name", toc + "/w:name", w + "val", "TOC 1"),
            ("TOC renamed id", toc, w + "styleId", "TOC11"),
            ("TOC line spacing", toc + "/w:pPr/w:spacing", w + "line", "360"),
        )
        for label, selector, attribute, value in cases:
            with self.subTest(mutation=label):
                root = ET.fromstring(after["word/styles.xml"])
                target = root.find(selector, NS)
                self.assertIsNotNone(target)
                target.set(attribute, value)
                tampered = {**after, "word/styles.xml": ET.tostring(root, encoding="utf-8")}
                with self.assertRaises(AssertionError):
                    self.assert_legacy_parts_match(before, tampered)
        with self.subTest(mutation="additional child in an otherwise reviewed TOC style"):
            root = ET.fromstring(after["word/styles.xml"])
            root.find(toc, NS).append(ET.Element(w + "qFormat"))
            tampered = {**after, "word/styles.xml": ET.tostring(root, encoding="utf-8")}
            with self.assertRaises(AssertionError):
                self.assert_legacy_parts_match(before, tampered)

    @unittest.skipUnless(BASELINE and Path(BASELINE).is_file(), "Set RUANZHU_MANUAL_BASELINE for HTML compatibility comparison")
    def test_legacy_html_equal_original(self):
        original = load_generator(Path(BASELINE), "manual_generator_baseline_html")
        for spec in ({"metadata": dict(META)}, generator.sample_spec(), legacy_spec()):
            with self.subTest(spec=spec["metadata"]["system_name"]):
                old_file, new_file = self.work / "old.html", self.work / "new.html"
                original.generate_html(copy.deepcopy(spec), old_file, self.work)
                generator.generate_html(copy.deepcopy(spec), new_file, self.work)
                self.assertEqual(old_file.read_bytes(), new_file.read_bytes())

    def test_library_outline_order_and_no_legacy_append(self):
        filename, document = self.generate(library_spec())
        self.assertEqual(headings(document), [
            ("Heading 1", "1 借还办理"), ("Heading 2", "1.1 借出登记"),
            ("Heading 3", "1.1.1 核对状态"), ("Heading 2", "1.2 归还登记"),
            ("Heading 1", "2 馆藏查询"), ("Heading 2", "2.1 筛选记录"),
        ])
        text = visible_text(package_parts(filename))
        self.assertNotIn("不得追加", text)
        self.assertNotIn("【需补充：", text)
        for heading in LEGACY_HEADINGS:
            self.assertNotIn(heading, text)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.inline_shapes), 4)

    def test_inspection_outline_and_sections_only_parents(self):
        filename, document = self.generate(inspection_spec())
        self.assertEqual(headings(document), [
            ("Heading 1", "1 异常闭环"), ("Heading 2", "1.1 指派复核"),
            ("Heading 3", "1.1.1 复核结论"), ("Heading 1", "2 巡检登记"),
            ("Heading 1", "3 班次交接"),
        ])
        text = visible_text(package_parts(filename))
        for heading in LEGACY_HEADINGS:
            self.assertNotIn(heading, text)
        self.assertEqual(len(document.inline_shapes), 2)
        self.assertEqual(len(document.tables), 1)

    def test_blocks_precede_sections_and_preserve_interleaving(self):
        filename, _ = self.generate(library_spec())
        body = ET.fromstring(package_parts(filename)["word/document.xml"]).find("w:body", NS)
        texts = ["".join(e.text or "" for e in child.findall(".//w:t", NS)) for child in body]
        expected = [
            "1 借还办理", "借还办理的内部测试说明。", "图 1-1 办理总览（内部测试图）",
            "1.1 借出登记", "1. 选择内部测试读者。", "2. 确认内部测试记录。",
            "- 保留当前选择。", "- 核对状态提示。", "1.1.1 核对状态", "三级标题下的内部测试说明。",
            "项目数量确认比率备注待借2True1.5已还0False0.0", "图 1-2 状态局部（内部测试图）",
            "1.2 归还登记", "图 1-3 归还示意（内部测试图）", "2 馆藏查询", "2.1 筛选记录",
        ]
        positions = [texts.index(text) for text in expected]
        self.assertEqual(positions, sorted(positions))

    def test_figure_numbers_accumulate_per_chapter(self):
        _, document = self.generate(library_spec())
        self.assertEqual([p.text for p in document.paragraphs if p.style.name == "图注"], [
            "图 1-1 办理总览（内部测试图）", "图 1-2 状态局部（内部测试图）",
            "图 1-3 归还示意（内部测试图）", "图 2-1 查询示意（内部测试图）",
        ])

    def test_lists_and_paragraphs_use_existing_body_style(self):
        _, document = self.generate(library_spec())
        texts = {p.text: p for p in document.paragraphs}
        for text in ("借还办理的内部测试说明。", "1. 选择内部测试读者。", "2. 确认内部测试记录。",
                     "- 保留当前选择。", "- 核对状态提示。"):
            self.assertEqual(texts[text].style.name, "正文段落")
            fonts = texts[text].runs[0]._r.rPr.rFonts
            self.assertEqual(fonts.get("{" + NS["w"] + "}eastAsia"), "宋体")
            self.assertEqual(fonts.get("{" + NS["w"] + "}ascii"), "Times New Roman")

    def test_table_scalars_and_three_line_borders(self):
        filename, document = self.generate(library_spec())
        self.assertEqual([[cell.text for cell in row.cells] for row in document.tables[0].rows], [
            ["项目", "数量", "确认", "比率", "备注"], ["待借", "2", "True", "1.5", ""], ["已还", "0", "False", "0.0", ""],
        ])
        table = ET.fromstring(package_parts(filename)["word/document.xml"]).find(".//w:tbl", NS)
        w = "{" + NS["w"] + "}"
        for row_index, row in enumerate(table.findall("w:tr", NS)):
            for cell in row.findall("w:tc", NS):
                borders = cell.find("w:tcPr/w:tcBorders", NS)
                for side in ("left", "right", "insideH", "insideV"):
                    self.assertEqual(borders.find("w:" + side, NS).get(w + "val"), "nil")
                for side in ("top", "bottom"):
                    edge = borders.find("w:" + side, NS)
                    expected = 12 if (row_index == 0 and side == "top") or (row_index == 2 and side == "bottom") else 8 if row_index == 0 and side == "bottom" else None
                    self.assertEqual(edge.get(w + "val"), "single" if expected else "nil")
                    if expected:
                        self.assertEqual(edge.get(w + "sz"), str(expected))

    def test_header_only_table_and_blank_scalar_cells_are_valid(self):
        for rows in ([], [[None, ""]]):
            with self.subTest(rows=rows):
                _, document = self.generate(self.block_spec({"type": "table", "headers": ["", "标记"], "rows": rows}))
                self.assertEqual(len(document.tables), 1)
                self.assertEqual(len(document.tables[0].rows), 1 + len(rows))
                self.assertEqual(document.tables[0].cell(0, 1).text, "标记")

    def test_inline_image_width_ratio_and_centering(self):
        from docx.shared import Cm
        filename, document = self.generate(library_spec())
        for shape, width_cm, ratio in zip(document.inline_shapes, (14.65, 5.08, 14.65, 14.65), (2, 0.5, 2, 2)):
            self.assertEqual(shape.width, Cm(width_cm))
            self.assertAlmostEqual(shape.width / shape.height, ratio, places=5)
        root = ET.fromstring(package_parts(filename)["word/document.xml"])
        self.assertEqual(len(root.findall(".//wp:inline", NS)), 4)
        self.assertEqual(root.findall(".//wp:anchor", NS), [])
        self.assertEqual([lock.get("noChangeAspect") for lock in root.findall(".//a:graphicFrameLocks", NS)], ["1"] * 4)
        for p in root.findall(".//w:p", NS):
            if p.find(".//wp:inline", NS) is not None:
                self.assertEqual(p.find("w:pPr/w:jc", NS).get("{" + NS["w"] + "}val"), "center")

    def test_toc_styles_are_native_lowercase_with_fixed_typography(self):
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        for spec in ({"metadata": dict(META)}, minimal_spec()):
            with self.subTest(mode="chapters" if "chapters" in spec else "legacy"):
                _, document = self.generate(spec)
                w = "{" + NS["w"] + "}"
                self.assertEqual(document.styles["Normal"].paragraph_format.first_line_indent, Pt(24))
                for level, left_pt in ((1, 0), (2, 24), (3, 48)):
                    name = f"toc {level}"
                    matches = [style for style in document.styles if style.name.lower() == name]
                    self.assertEqual(len(matches), 1, name)
                    style = document.styles[name]
                    self.assertEqual(style.name, name)
                    self.assertEqual(style.element.find(w + "name").get(w + "val"), name)
                    self.assertEqual(style.style_id, f"toc{level}")
                    self.assertTrue(style.builtin)
                    self.assertIsNone(style.element.get(w + "customStyle"))
                    self.assertEqual(style.type, WD_STYLE_TYPE.PARAGRAPH)
                    self.assertEqual(style.base_style.style_id, "Normal")
                    self.assertEqual(style.font.size, Pt(12))
                    self.assertFalse(style.font.bold)
                    self.assertEqual(str(style.font.color.rgb), "000000")
                    fonts = style.element.rPr.rFonts
                    self.assertEqual(fonts.get(w + "eastAsia"), "宋体")
                    for script in ("ascii", "hAnsi", "cs"):
                        self.assertEqual(fonts.get(w + script), "Times New Roman")
                    paragraph_format = style.paragraph_format
                    self.assertEqual(paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)
                    self.assertEqual(paragraph_format.left_indent, Pt(left_pt))
                    self.assertEqual(paragraph_format.first_line_indent, Pt(0))
                    self.assertEqual(paragraph_format.space_before, Pt(0))
                    self.assertEqual(paragraph_format.space_after, Pt(0))
                    self.assertEqual(paragraph_format.line_spacing, 1.0)
                    indent = style.element.pPr.find(w + "ind")
                    self.assertEqual(indent.get(w + "left"), str(left_pt * 20))
                    self.assertEqual(indent.get(w + "firstLine"), "0")

    def test_image_paragraphs_explicitly_clear_first_line_indent(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        for spec in (legacy_spec(), library_spec()):
            with self.subTest(mode="chapters" if "chapters" in spec else "legacy"):
                _, document = self.generate(spec)
                images = [p for p in document.paragraphs if p._p.find(".//{" + NS["wp"] + "}inline") is not None]
                self.assertGreater(len(images), 0)
                self.assertEqual(len(images), len(document.inline_shapes))
                self.assertEqual(document.styles["Normal"].paragraph_format.first_line_indent, Pt(24))
                for image in images:
                    self.assertEqual(image.style.name, "Normal")
                    self.assertEqual(image.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                    self.assertIsNotNone(image.paragraph_format.first_line_indent)
                    self.assertEqual(image.paragraph_format.first_line_indent, Pt(0))
                    self.assertIsNone(image.paragraph_format.left_indent)
                    self.assertIsNone(image.paragraph_format.right_indent)
                    w = "{" + NS["w"] + "}"
                    indent = image._p.pPr.find(w + "ind")
                    self.assertEqual(dict(indent.attrib), {w + "firstLine": "0"})
                    self.assertEqual(image._p.pPr.find(w + "jc").get(w + "val"), "center")

    def test_absolute_image_paths_and_sources_are_not_visible(self):
        spec = minimal_spec()
        internal = str(self.work / "images" / "small.png")
        spec["sources"] = [r"C:\private\database.sqlite", "source.py", internal]
        spec["chapters"][0]["blocks"].append({"type": "image", "image": internal, "caption": "内部测试图"})
        filename, document = self.generate(spec)
        text = visible_text(package_parts(filename))
        for source in spec["sources"]:
            self.assertNotIn(source, text)
        self.assertNotIn("small.png", text)
        self.assertEqual(len(document.inline_shapes), 1)

    def test_custom_styles_sections_headers_and_settings_match_legacy(self):
        spec = library_spec()
        old_spec = legacy_spec()
        old_spec["metadata"] = spec["metadata"]
        custom_file, custom_doc = self.generate(spec, "custom.docx")
        legacy_file, _ = self.generate(old_spec, "legacy.docx")
        custom, legacy = package_parts(custom_file), package_parts(legacy_file)
        for name in legacy:
            if name in {"word/styles.xml", "word/stylesWithEffects.xml", "word/settings.xml", "word/fontTable.xml", "word/numbering.xml"} or name.startswith(("word/header", "word/footer")):
                self.assertEqual(custom[name], legacy[name], name)
        sections = lambda parts: [ET.tostring(section) for section in ET.fromstring(parts["word/document.xml"]).findall(".//w:sectPr", NS)]
        self.assertEqual(sections(custom), sections(legacy))
        self.assertEqual(len(custom_doc.sections), 3)

    def test_fixed_three_section_page_field_structure(self):
        from docx.shared import Cm
        filename, _ = self.generate(minimal_spec())
        parts = package_parts(filename)
        root = ET.fromstring(parts["word/document.xml"])
        sections = root.findall(".//w:sectPr", NS)
        self.assertEqual(len(sections), 3)
        w, r = "{" + NS["w"] + "}", "{" + NS["r"] + "}"
        for section in sections:
            size, margins = section.find("w:pgSz", NS), section.find("w:pgMar", NS)
            self.assertEqual((size.get(w + "w"), size.get(w + "h")), (str(Cm(21).twips), str(Cm(29.7).twips)))
            for key, cm in {"left": 3.17, "right": 3.17, "top": 2.54, "bottom": 2.54, "header": 1.5, "footer": 1.75}.items():
                self.assertEqual(margins.get(w + key), str(Cm(cm).twips))
        self.assertIsNone(sections[0].find("w:headerReference", NS))
        self.assertIsNone(sections[0].find("w:footerReference", NS))
        self.assertIsNone(sections[0].find("w:pgNumType", NS))
        relationships = {link.get("Id"): link.get("Target") for link in ET.fromstring(parts["word/_rels/document.xml.rels"])}
        for index, section in enumerate(sections[1:], start=1):
            self.assertEqual(section.find("w:pgNumType", NS).get(w + "start"), "1")
            for kind, expected in (("header", [] if index == 1 else ["PAGE", "SECTIONPAGES"]), ("footer", ["PAGE"] if index == 1 else [])):
                reference = section.find("w:" + kind + "Reference", NS)
                part = ET.fromstring(parts["word/" + relationships[reference.get(r + "id")]])
                self.assertEqual([e.text for e in part.findall(".//w:instrText", NS)], expected)
                if kind == "header":
                    border = part.find(".//w:pBdr/w:bottom", NS)
                    self.assertEqual((border.get(w + "color"), border.get(w + "sz")), ("A6A6A6", "4"))
        self.assertEqual([e.text for e in root.findall(".//w:instrText", NS)], ['TOC \\o "1-3" \\h \\z \\u'])
        update = ET.fromstring(parts["word/settings.xml"]).find("w:updateFields", NS)
        self.assertEqual(update.get(w + "val"), "true")

    def test_no_chapter_cap_or_artificial_page_breaks(self):
        spec = minimal_spec()
        spec["chapters"] = [{"title": f"业务项 {i}", "blocks": [paragraph()]} for i in range(1, 102)]
        filename, document = self.generate(spec)
        self.assertEqual(len(headings(document)), 101)
        self.assertEqual(headings(document)[-1], ("Heading 1", "101 业务项 101"))
        root = ET.fromstring(package_parts(filename)["word/document.xml"])
        self.assertEqual(len(root.findall(".//w:sectPr", NS)), 3)
        self.assertEqual(root.findall(".//w:br[@w:type='page']", NS), [])
        self.assertEqual(root.findall(".//w:pPr/w:pageBreakBefore", NS), [])

    def test_custom_mode_ignores_generic_fillers_without_mutating_input(self):
        spec = library_spec()
        spec["_allow_generic_defaults"] = True
        before = copy.deepcopy(spec)
        filename, _ = self.generate(spec)
        self.assertEqual(spec, before)
        text = visible_text(package_parts(filename))
        self.assertNotIn("不得追加", text)
        self.assertNotIn("【需补充：", text)
        for heading in LEGACY_HEADINGS:
            self.assertNotIn(heading, text)

    def test_invalid_chapters_containers_fail_before_document_creation(self):
        for chapters in (None, [], {}, "章节", False, (minimal_spec()["chapters"][0],)):
            with self.subTest(chapters=chapters):
                self.reject({"metadata": dict(META), "chapters": chapters})

    def test_invalid_nodes_titles_and_children_fail(self):
        nodes = [
            None, "章节", [], {}, {"blocks": [paragraph()]},
            *({"title": value, "blocks": [paragraph()]} for value in (None, "", " \t\n", 3, False)),
            {"title": "空章"}, {"title": "空章", "blocks": [], "sections": []},
            *({"title": "业务", "blocks": value} for value in (None, {}, "正文", (paragraph(),))),
            *({"title": "业务", "blocks": [paragraph()], "sections": value} for value in (None, {}, "小节", ())),
            {"title": "业务", "blocks": [paragraph()], "sections": [{"title": "空小节"}]},
        ]
        for node in nodes:
            with self.subTest(node=node):
                self.reject({"metadata": dict(META), "chapters": [node]})

    def test_unknown_node_and_block_fields_are_rejected(self):
        for key in ("unknown", "style", "page_break", "path"):
            spec = minimal_spec()
            spec["chapters"][0][key] = "不允许静默丢弃"
            with self.subTest(node_field=key):
                self.reject(spec, "未知字段")
        blocks = [
            paragraph(), {"type": "steps", "items": ["步骤"]}, {"type": "bullets", "items": ["提示"]},
            {"type": "table", "headers": ["字段"], "rows": [["值"]]},
            {"type": "image", "image": "images/wide.png", "caption": "内部测试图"},
        ]
        for block in blocks:
            with self.subTest(block_type=block["type"]):
                self.reject(self.block_spec({**block, "unknown": "不得忽略"}), "未知字段")

    def test_fourth_heading_level_is_rejected(self):
        node = minimal_spec()["chapters"][0]
        for _ in range(3):
            node = {"title": "父节点", "sections": [node]}
        self.reject({"metadata": dict(META), "chapters": [node]}, "最多为 3 级")

    def test_unknown_missing_and_nonstring_block_types_fail(self):
        blocks = [None, "正文", [], {}, *({"type": kind} for kind in ("page_break", "unknown", "", None, 1, [], {}))]
        for block in blocks:
            with self.subTest(block=block):
                self.reject(self.block_spec(block))

    def test_paragraph_requires_nonblank_string(self):
        for block in [{"type": "paragraph"}, *({"type": "paragraph", "text": value} for value in (None, "", " \n", 0, False, {}, []))]:
            with self.subTest(block=block):
                self.reject(self.block_spec(block))

    def test_steps_and_bullets_reject_illegal_items(self):
        for kind in ("steps", "bullets"):
            for value in (None, [], "步骤", {}, ("步骤",), [""], [" \t"], ["有效", 3], [None], [False], [{}], [[]]):
                with self.subTest(kind=kind, items=value):
                    self.reject(self.block_spec({"type": kind, "items": value}))
            self.reject(self.block_spec({"type": kind}))

    def test_tables_reject_bad_headers_rows_and_non_scalars(self):
        valid = {"type": "table", "headers": ["字段", "值"], "rows": [["甲", 2]]}
        blocks = [
            {"type": "table", "rows": []}, {"type": "table", "headers": ["字段"]},
            *({**valid, "headers": value} for value in (None, [], "字段", [1], [None], ("字段", "值"))),
            *({**valid, "rows": value} for value in (None, "行", {}, (), [None], ["甲乙"], [{"字段": "值"}], [("甲", 2)], [[]], [["甲"]], [["甲", 2, "多余"]])),
            *({**valid, "rows": [["甲", value]]} for value in ({}, [], (1,), {1})),
        ]
        for block in blocks:
            with self.subTest(block=block):
                self.reject(self.block_spec(block))

    def test_images_require_path_and_caption_strings(self):
        valid = {"type": "image", "image": "images/wide.png", "caption": "内部测试图"}
        for field in ("image", "caption"):
            block = {key: value for key, value in valid.items() if key != field}
            self.reject(self.block_spec(block))
            for value in (None, "", " \t", 1, False, {}, []):
                with self.subTest(field=field, value=value):
                    self.reject(self.block_spec({**valid, field: value}))

    def test_missing_image_fails_before_write_without_leaking_path(self):
        private_path = str(self.work / "private-source" / "absent.png")
        spec = minimal_spec()
        spec["chapters"][0]["blocks"] += [
            {"type": "image", "image": "images/wide.png", "caption": "已读取的内部测试图"},
            {"type": "image", "image": private_path, "caption": "缺失的内部测试图"},
        ]
        message = self.reject(spec, "截图缺失", preflight=False)
        self.assertNotIn(private_path, message)
        self.assertNotIn("absent.png", message)
        self.assertNotIn("private-source", message)

    def test_corrupt_or_directory_image_fails_without_leaking_path(self):
        bad_image = self.work / "private-invalid.png"
        bad_image.write_bytes(b"not an image")
        for image in (bad_image, self.work / "images"):
            with self.subTest(kind=image.name):
                message = self.reject(self.block_spec({"type": "image", "image": str(image), "caption": "内部测试图"}), "截图文件无法读取", preflight=False)
                self.assertNotIn(str(image), message)

    def test_paths_in_all_custom_visible_fields_are_rejected(self):
        private = r"C:\private\source.py"
        specs = []
        for field in META:
            spec = minimal_spec()
            spec["metadata"][field] = private
            specs.append(spec)
        spec = minimal_spec()
        spec["chapters"][0]["title"] = private
        specs.append(spec)
        blocks = [
            paragraph(private), {"type": "steps", "items": [private]}, {"type": "bullets", "items": [private]},
            {"type": "table", "headers": [private], "rows": [["值"]]},
            {"type": "table", "headers": ["字段"], "rows": [[private]]},
            {"type": "image", "image": "images/wide.png", "caption": private},
        ]
        specs.extend(self.block_spec(block) for block in blocks)
        for index, spec in enumerate(specs):
            with self.subTest(visible_field=index):
                message = self.reject(spec, "本机文件名或路径")
                self.assertNotIn(private, message)

    def test_existing_filename_and_path_detection_is_retained(self):
        for value in ("src/internal.py", r"\\server\share\database.sqlite", "file:///C:/private/view.png", "README.md"):
            with self.subTest(value=value):
                self.reject(self.block_spec(paragraph(value)), "本机文件名或路径")
                self.reject({"metadata": dict(META), "summary": value}, "本机文件名或路径")

    def test_validation_failures_do_not_overwrite_existing_docx(self):
        output = self.work / "protected.docx"
        sentinel = b"internal test sentinel - must remain unchanged"
        output.write_bytes(sentinel)
        specs = [
            {"chapters": []}, self.block_spec(paragraph(r"C:\private\source.py")),
            self.block_spec({"type": "image", "image": "private-absent.png", "caption": "内部测试图"}),
        ]
        for spec in specs:
            with self.subTest(spec=spec):
                with self.assertRaises(SystemExit):
                    generator.generate_docx(spec, output, self.work)
                self.assertEqual(output.read_bytes(), sentinel)

    def test_html_custom_mode_rejects_before_any_output_or_pdf(self):
        for chapters in (None, [], {}, minimal_spec()["chapters"]):
            with self.subTest(chapters=chapters):
                output, pdf = self.work / "unwritten-html" / "manual.html", self.work / "unwritten-pdf" / "manual.pdf"
                with self.assertRaises(SystemExit) as caught:
                    generator.generate_html({"chapters": chapters}, output, self.work, pdf)
                self.assertIn("HTML", str(caught.exception))
                self.assertIn("--backend docx", str(caught.exception))
                self.assertFalse(output.parent.exists())
                self.assertFalse(pdf.parent.exists())
        output, pdf = self.work / "protected.html", self.work / "protected.pdf"
        for filename in (output, pdf):
            filename.write_bytes(b"preserved internal sentinel")
        with self.assertRaises(SystemExit):
            generator.generate_html(minimal_spec(), output, self.work, pdf)
        for filename in (output, pdf):
            self.assertEqual(filename.read_bytes(), b"preserved internal sentinel")

    def test_cli_custom_docx_resolves_images_relative_to_input(self):
        from docx import Document
        project = self.work / "project"
        make_images(project)
        spec_file = project / "project.json"
        spec_file.write_text(json.dumps(library_spec(), ensure_ascii=False), encoding="utf-8")
        other_cwd = self.work / "other-cwd"
        other_cwd.mkdir()
        output = other_cwd / "custom.docx"
        result = self.cli("--input", Path("..") / "project" / "project.json", "--output", output, "--backend", "docx", cwd=other_cwd)
        self.assertEqual(result.returncode, 0, result.stderr)
        document = Document(output)
        self.assertEqual(len(document.inline_shapes), 4)
        self.assertEqual(len(headings(document)), 6)
        self.assertEqual(headings(document)[0], ("Heading 1", "1 借还办理"))

    def test_cli_errors_are_clean_and_do_not_write(self):
        cases = [
            ({"chapters": []}, "docx", "非空列表"),
            (self.block_spec({"type": "page_break"}), "docx", "paragraph"),
            (self.block_spec({"type": "table", "headers": ["列"], "rows": [["多", "余"]]}), "docx", "列数"),
            (self.block_spec({"type": "image", "image": "private-absent.png", "caption": "内部测试图"}), "docx", "截图缺失"),
            (self.block_spec(paragraph(r"C:\private\source.py")), "docx", "本机文件名或路径"),
            (minimal_spec(), "html", "--backend docx"),
        ]
        for index, (spec, backend, message) in enumerate(cases):
            with self.subTest(case=index):
                input_file = self.work / f"invalid-{index}.json"
                input_file.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
                output = self.work / "unwritten-cli" / f"manual.{backend}"
                pdf = self.work / "unwritten-cli-pdf" / "manual.pdf"
                arguments = ["--input", input_file, "--output", output, "--backend", backend]
                if backend == "html":
                    arguments += ["--pdf", pdf]
                result = self.cli(*arguments)
                self.assertEqual(result.returncode, 1, result.stderr)
                self.assertIn(message, result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertNotIn("private-absent.png", result.stderr)
                self.assertNotIn(r"C:\private", result.stderr)
                self.assertEqual(result.stdout, "")
                self.assertFalse(output.parent.exists())
                self.assertFalse(pdf.parent.exists())

    def test_cli_sample_default_output_generic_flag_and_legacy_html(self):
        sample = self.work / "sample.json"
        result = self.cli("--write-sample", sample)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(sample.read_text(encoding="utf-8")), generator.sample_spec())
        spec_file = self.work / "minimal-legacy.json"
        spec_file.write_text(json.dumps({"metadata": dict(META)}, ensure_ascii=False), encoding="utf-8")
        result = self.cli("--input", spec_file, "--allow-generic-defaults")
        self.assertEqual(result.returncode, 0, result.stderr)
        default_output = self.work / f"{META['system_name']}{META['version']}{META['doc_type']}.docx"
        text = visible_text(package_parts(default_output))
        self.assertNotIn("【需补充：", text)
        self.assertIn("8 总结", text)
        html_output = self.work / "legacy.html"
        result = self.cli("--input", spec_file, "--backend", "html", "--output", html_output)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("【需补充：编写目的】", html_output.read_text(encoding="utf-8"))
        self.assertIn("<h1>8 总结</h1>", html_output.read_text(encoding="utf-8"))

    def test_public_artifacts_default_still_rejects_known_filenames(self):
        for filename in ("applet.py", "runtime.yaml", "report.csv"):
            with self.subTest(filename=filename):
                self.assertTrue(generator.contains_local_artifact(filename))
                with self.assertRaises(SystemExit):
                    generator.validate_public_content({"nested": [{"text": filename}]})
                self.reject(self.block_spec(paragraph(f"打开 “{filename}”。")), "本机文件名或路径")

    def test_public_artifacts_allow_exact_literals_in_all_requested_docx_blocks(self):
        records = public_artifact_records()
        quoted_names = "、".join(f"“{record['value']}”" for record in records)
        body_text = "已核对的交付入口：" + quoted_names
        spec = minimal_spec()
        spec["public_artifacts"] = records
        spec["chapters"][0]["blocks"] = [
            paragraph(body_text),
            {"type": "steps", "items": ["使用 " + quoted_names]},
            {"type": "table", "headers": ["交付入口", "操作说明"],
             "rows": [[record["value"], "按已核对入口办理"] for record in records]},
            {"type": "image", "image": "images/wide.png", "caption": "交付入口示意：" + quoted_names},
        ]
        unchanged = copy.deepcopy(spec)
        filename, document = self.generate(spec)
        self.assertEqual(spec, unchanged)
        self.assertEqual(set(generator.public_artifact_literals(spec)), {record["value"] for record in records})
        paragraph_texts = [p.text for p in document.paragraphs]
        self.assertIn(body_text, paragraph_texts)
        self.assertIn("1. 使用 " + quoted_names, paragraph_texts)
        table_values = [row.cells[0].text for row in document.tables[0].rows[1:]]
        self.assertEqual(table_values, [record["value"] for record in records])
        self.assertEqual([p.text for p in document.paragraphs if p.style.name == "图注"],
                         ["图 1-1 交付入口示意：" + quoted_names])
        text = visible_text(package_parts(filename))
        for record in records:
            with self.subTest(literal=record["value"]):
                self.assertIn(record["value"], text)
                self.assertNotIn(record["reason"], text)
                self.assertNotIn(record["evidence"], text)
        self.assertNotIn("public_artifacts", text)

    def test_public_artifacts_do_not_allow_other_names_or_private_paths(self):
        cases = (
            ("unregistered filename", "unlisted.csv"),
            ("basename of a registered relative path", "runtime.yaml"),
            ("extra prefix directory", "src/applet.py"),
            ("Windows relative prefix directory", r"src\applet.py"),
            ("longer filename", "myapplet.py"),
            ("longer extension", "applet.py.csv"),
            ("absolute Windows path", r"C:\private\applet.py"),
            ("UNC path", r"\\server\share\applet.py"),
            ("file URL", "file:///C:/private/applet.py"),
        )
        for label, value in cases:
            with self.subTest(case=label):
                spec = self.block_spec(paragraph(f"打开 “{value}”。"))
                spec["public_artifacts"] = public_artifact_records()
                message = self.reject(spec, "本机文件名或路径")
                self.assertNotIn(value, message)

    def test_public_artifacts_invalid_registry_fails_before_write_or_overwrite(self):
        record = public_artifact_records()[0]
        cases = [
            ("null list", None),
            ("object instead of list", record),
            ("string instead of record", ["applet.py"]),
            ("missing evidence", [{key: value for key, value in record.items() if key != "evidence"}]),
            ("unknown record field", [{**record, "extra": "must not be ignored"}]),
            ("nonstring evidence", [{**record, "evidence": False}]),
            ("duplicate literal", [record, {**record, "reason": "另一份内部核对说明"}]),
        ]
        cases.extend((f"blank {field}", [{**record, field: " \t"}]) for field in ("value", "reason", "evidence"))
        cases.extend((label, [{**record, "value": value}]) for label, value in (
            ("absolute value", r"C:\private\applet.py"),
            ("UNC value", r"\\server\share\applet.py"),
            ("file URL value", "file:///C:/private/applet.py"),
            ("parent traversal", "../applet.py"),
            ("Windows nested traversal", r"config\..\applet.py"),
            ("star wildcard", "*.py"),
            ("question wildcard", "config/?.yaml"),
        ))
        sentinel = b"internal sentinel: no overwrite on invalid public registry"
        for backend, render in (("docx", generator.generate_docx), ("html", generator.generate_html)):
            for label, registry in cases:
                with self.subTest(backend=backend, case=label):
                    spec = minimal_spec() if backend == "docx" else {"metadata": dict(META)}
                    spec["public_artifacts"] = copy.deepcopy(registry)
                    missing_output = self.work / "unwritten-public-registry" / f"manual.{backend}"
                    existing_output = self.work / f"protected-public.{backend}"
                    existing_output.write_bytes(sentinel)
                    with mock.patch("docx.Document") as document_factory:
                        for output in (missing_output, existing_output):
                            with self.assertRaises(SystemExit) as caught:
                                render(spec, output, self.work)
                            self.assertIn("public_artifacts", str(caught.exception))
                        document_factory.assert_not_called()
                    self.assertFalse(missing_output.parent.exists())
                    self.assertEqual(existing_output.read_bytes(), sentinel)

    def test_public_artifacts_legacy_html_optional_behavior(self):
        records = public_artifact_records()
        quoted_names = "、".join(f"“{record['value']}”" for record in records)
        spec = legacy_spec()
        spec["public_artifacts"] = records
        spec["overview"]["purpose"] = "已核对的交付入口：" + quoted_names
        spec["environment"]["startup_steps"] = ["按登记入口使用 " + quoted_names]
        spec["modules"][0]["caption"] = "已登记文件：" + quoted_names
        spec["data_outputs"][0]["description"] = "交付入口：" + quoted_names
        output = self.work / "public-artifacts.html"
        generator.generate_html(spec, output, self.work)
        html_text = output.read_text(encoding="utf-8")
        for record in records:
            with self.subTest(literal=record["value"]):
                self.assertIn(record["value"], html_text)
                self.assertNotIn(record["reason"], html_text)
                self.assertNotIn(record["evidence"], html_text)
        without_registry = copy.deepcopy(spec)
        del without_registry["public_artifacts"]
        before = output.read_bytes()
        with self.assertRaises(SystemExit):
            generator.generate_html(without_registry, output, self.work)
        self.assertEqual(output.read_bytes(), before)
        # An unused review record must not change a legacy HTML byte. The existing
        # original-generator HTML comparison separately checks the same legacy spec.
        clean = legacy_spec()
        plain_output = self.work / "no-public-artifacts.html"
        generator.generate_html(clean, plain_output, self.work)
        for registry in ([], records):
            with self.subTest(unused_registry_size=len(registry)):
                generator.generate_html({**clean, "public_artifacts": registry}, output, self.work)
                self.assertEqual(output.read_bytes(), plain_output.read_bytes())


if __name__ == "__main__":
    unittest.main()
